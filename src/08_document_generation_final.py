import re
from docx import Document
from .common import get,get_json,OUTPUT,safe
def run(process,persona,country,conversation_id,**kwargs):
 sid=safe(conversation_id); text=get(f'{process}_{persona}_{country}_{sid}_EUTM_response.txt'); subs=get_json(f'{process}_{persona}_{country}_{sid}_sub_bp_tasks.json')
 for k,v in subs.items(): text=text.replace(k+'_task',v or '')
 text='\n'.join(x for x in text.splitlines() if '_task' not in x and '_scenario' not in x); doc=Document(); doc.add_heading(f'{process} for Location {country}',0); lines=text.splitlines(); i=0
 while i<len(lines):
  s=lines[i].strip()
  if s.startswith('# '): doc.add_heading(s[2:],1)
  elif s.startswith('## '): doc.add_heading(s[3:],2)
  elif s.startswith('### '): doc.add_heading(s[4:],3)
  elif s.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'): rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i+=1
   rows=[r for r in rows if not all(re.fullmatch(r':?-{3,}:?',c or '') for c in r)]
   if rows:
    table=doc.add_table(rows=1,cols=max(map(len,rows))); table.style='Table Grid'
    for j,c in enumerate(rows[0]): table.rows[0].cells[j].text=c
    for row in rows[1:]:
     cells=table.add_row().cells
     for j,c in enumerate(row): cells[j].text=c
   continue
  elif s: doc.add_paragraph(s)
  i+=1
 name=f'{process}_{persona}_for_{country}_{sid}_EUTM_Document.docx'; doc.save(OUTPUT/name); return {'document':str(OUTPUT/name)}
